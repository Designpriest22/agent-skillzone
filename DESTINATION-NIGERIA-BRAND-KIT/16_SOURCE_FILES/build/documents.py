"""Destination Nigeria document template library → 09_DOCUMENT_TEMPLATES (DOCX; PDFs via LibreOffice).

Rules baked in: no invented addresses, phone numbers, bank details, registration numbers, seals,
affiliations, partners, statistics or dates. Every such item is a bracketed content-control placeholder.
"""
import json
import subprocess
from docx.shared import Mm, Pt
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from brand import *
from docbase import *
from brand import CONTACT as _CONTACT

D = KIT / "09_DOCUMENT_TEMPLATES"
saved = []


def save(doc, rel):
    p = D / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(p)
    saved.append(p)
    return p


CONTACT_LINES = [_CONTACT["phone"], _CONTACT["email"], _CONTACT["web"]]


# ====================================================================== charts & images used inside documents
def crop_to(src, out, w_mm, h_mm, focus=(0.5, 0.5), px_per_mm=11.8):
    """Crop a library photograph to an exact frame (default ≈300 ppi) so Word places it without distortion."""
    from PIL import Image
    im = Image.open(src).convert("RGB")
    W, H = im.size
    r = w_mm / h_mm
    if W / H > r:
        cw, ch = int(H * r), H
    else:
        cw, ch = W, int(W / r)
    x0 = int(min(max(focus[0] * W - cw / 2, 0), W - cw))
    y0 = int(min(max(focus[1] * H - ch / 2, 0), H - ch))
    im = im.crop((x0, y0, x0 + cw, y0 + ch))
    tw = int(w_mm * px_per_mm)
    if im.width > tw:
        im = im.resize((tw, int(tw / r)), Image.LANCZOS)
    im.save(out, quality=90)
    return out


def make_assets():
    """Document artwork: cropped library photographs, plate pages and an editorial chart. No invented imagery."""
    import re
    A = KIT / "16_SOURCE_FILES/build/_doc_assets"
    A.mkdir(exist_ok=True)
    jobs = []
    # chart: editorial, Public Sans, Forest + Sage, labelled as illustrative
    W, H = 1600, 760
    vals_a, vals_b = [42, 55, 49, 63, 71, 78], [30, 34, 38, 41, 47, 52]
    labs = ["Q1 2025", "Q2", "Q3", "Q4", "Q1 2026", "Q2"]
    g = ""
    for k in range(0, 9):
        y = 620 - k * 70
        g += f'<path d="M140,{y} H1540" stroke="{C["haze"] if k else C["forest"]}" stroke-width="{1.5 if k else 2.5}"/>'
        g += f'<text x="120" y="{y + 8}" text-anchor="end" font-family="Public Sans" font-size="22" fill="{C["stone"]}">{k * 10}</text>'
    for i, (va, vb) in enumerate(zip(vals_a, vals_b)):
        x = 200 + i * 225
        g += f'<rect x="{x}" y="{620 - va * 7}" width="72" height="{va * 7}" fill="{C["forest"]}"/>'
        g += f'<rect x="{x + 80}" y="{620 - vb * 7}" width="72" height="{vb * 7}" fill="{C["sage"]}"/>'
        g += f'<text x="{x + 76}" y="660" text-anchor="middle" font-family="Public Sans" font-size="22" fill="{C["charcoal"]}">{labs[i]}</text>'
    g += (f'<text x="140" y="50" font-family="Public Sans" font-weight="600" font-size="20" letter-spacing="1.6" fill="{C["stone"]}">ILLUSTRATIVE DATA: REPLACE WITH VERIFIED FIGURES</text>'
          f'<rect x="1140" y="36" width="20" height="20" fill="{C["forest"]}"/><text x="1170" y="53" font-family="Public Sans" font-size="22" fill="{C["charcoal"]}">Series A</text>'
          f'<rect x="1320" y="36" width="20" height="20" fill="{C["sage"]}"/><text x="1350" y="53" font-family="Public Sans" font-size="22" fill="{C["charcoal"]}">Series B</text>')
    (A / "chart.svg").write_text(svg_doc(W, H, f'<rect width="{W}" height="{H}" fill="#fff"/>{g}', "Illustrative chart"))
    jobs.append(dict(src=str(A / "chart.svg"), out=str(A / "chart.png"), w=W, h=H, scale=1.5))
    # photographs cropped to their frames
    crop_to(photo("zuma-rock"), A / "cover_report.jpg", 210, 168, (0.5, 0.45))
    crop_to(photo("agbokim-waterfalls"), A / "photo_invest.jpg", 210, 92, (0.45, 0.35))
    crop_to(photo("idanre-hills"), A / "photo_body.jpg", 166, 100, (0.4, 0.32))
    crop_to(photo("idanre-hills"), A / "cover_workshop.jpg", 148.5, 210, (0.42, 0.4))
    # plate pages: proposal cover (Forest plate on Paper) and chapter divider (Paper plate on Forest)
    def plate_page(slug, ink, bg, w, h, x, y, pw, name, cap=None, capcol=None):
        pl = plate(slug, ink, True).read_text()
        vb = re.search(r'viewBox="([^"]+)"', pl).group(1)
        inner = pl[pl.index(">", pl.index("<svg")) + 1:pl.rindex("</svg>")]
        inner = re.sub(r"<title.*?</title>|<desc.*?</desc>", "", inner, flags=re.S)
        vw, vh = [float(v) for v in vb.split()[2:]]
        ph_ = pw * vh / vw
        body = (f'<rect width="{w}" height="{h}" fill="{bg}"/>'
                f'<svg x="{x}" y="{y}" width="{pw}" height="{ph_}" viewBox="{vb}">{inner}</svg>')
        if cap:
            body += f'<text x="{x}" y="{y + ph_ + 6}" font-family="Public Sans" font-size="3.1" fill="{capcol}">{cap}</text>'
        (A / f"{name}.svg").write_text(svg_doc(w, h, body, name))
        jobs.append(dict(src=str(A / f"{name}.svg"), out=str(A / f"{name}.jpg"), w=w * 6, h=h * 6, quality=92, scale=2))
    plate_page("zuma-rock", "forest", C["ivory"], 210, 297, 22, 30, 166, "cover_proposal",
               "Pl. 01   Zuma Rock, Niger State", C["stone"])
    plate_page("agbokim-waterfalls", "ivory", C["forest"], 210, 297, 96, 0, 114, "divider")
    # certificate: Paper with an engraver's double rule, nothing else
    cert = (f'<rect width="297" height="210" fill="{C["ivory"]}"/>'
            f'<rect x="10" y="10" width="277" height="190" fill="none" stroke="{C["forest"]}" stroke-width="0.6"/>'
            f'<rect x="12.2" y="12.2" width="272.6" height="185.6" fill="none" stroke="{C["forest"]}" stroke-width="0.2"/>')
    (A / "certificate_bg.svg").write_text(svg_doc(297, 210, cert, "Certificate ground"))
    jobs.append(dict(src=str(A / "certificate_bg.svg"), out=str(A / "certificate_bg.jpg"), w=297 * 6, h=210 * 6, quality=92, scale=1.5))
    json.dump(jobs, open(A / "jobs.json", "w"))
    subprocess.run(["node", str(KIT / "16_SOURCE_FILES/build/render.mjs"), str(A / "jobs.json")], check=True)
    return A


# ====================================================================== letterhead & letter
def letterhead(variant="standard", sample=False):
    d = new(margins=(46, 22, 30, 24), title="Destination Nigeria letterhead", subject="Official correspondence")
    s = d.sections[0]
    s.different_first_page_header_footer = True
    if variant == "standard":
        logo_header(s, 46, CONTACT_LINES, first=True)
    else:
        h = s.first_page_header
        h.is_linked_to_previous = False
        p = h.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Mm(58))
    small_header(s, "Continued")
    contact_footer(s, page_numbers=True)
    # first-page footer: contact line
    ff = s.first_page_footer
    ff.is_linked_to_previous = False
    fp = ff.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if variant == "ceremonial":
        run(fp, "   ".join(CONTACT_LINES), color=C["stone"], size=7.5)
    else:
        run(fp, "Destination Nigeria", bold=True, color=C["forest"], size=7.5)
        run(fp, "   Experience. Invest. Explore.", color=C["stone"], size=7.5)
    if not sample:
        P(d, "")
        return d
    # formal correspondence layout
    t = d.add_table(rows=1, cols=2)
    borders(t, "#FFFFFF", 0, outer=False)
    a, b = t.rows[0].cells
    pa = a.paragraphs[0]
    pa.style = d.styles["Small"]
    run(pa, "Our ref: ")
    field(pa, "[DN/XXX/YYYY/000]", "reference")
    pb = b.paragraphs[0]
    pb.style = d.styles["Small"]
    pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(pb, "Date: ")
    field(pb, "[DD Month YYYY]", "date")
    widths(t, [83, 83])
    P(d, "")
    for ph in ["[Recipient name]", "[Title / Position]", "[Organisation]", "[Address line 1]", "[City, State]"]:
        q = P(d, ph=ph)
        q.paragraph_format.space_after = Pt(0)
    P(d, "")
    P(d, ph="[Dear Sir / Madam / Name],")
    q = P(d, "", style="Heading 3")
    run(q, "RE: ")
    field(q, "[SUBJECT OF THE LETTER IN CAPITALS]", "subject", color=C["forest"], bold=True)
    for txt in ["[Opening paragraph — state the purpose of the letter in one or two sentences.]",
                "[Body — provide context, facts and any request. Use verified information only; cite documents by title and date.]",
                "[Closing paragraph — state the next step, the contact person and the response date if one is required.]"]:
        P(d, ph=txt)
    P(d, "Yours faithfully,")
    P(d, "")
    P(d, "")
    q = P(d, ph="[Name of signatory]")
    q.paragraph_format.space_after = Pt(0)
    q = P(d, ph="[Title / Position]")
    q.paragraph_format.space_after = Pt(0)
    P(d, "For: Destination Nigeria", style="Small")
    P(d, "")
    q = P(d, "Enclosures: ", style="Small")
    field(q, "[List or ‘None’]")
    q = P(d, "cc: ", style="Small")
    field(q, "[Names]")
    return d


# ====================================================================== memo
def memo():
    d = new(title="Internal memorandum")
    s = d.sections[0]
    logo_header(s, 40)
    contact_footer(s, "Internal memorandum")
    P(d, "Internal", style="Overline")
    P(d, "Memorandum", style="Title")
    field_table(d, [("Reference", "[DN/MEMO/YYYY/000]"), ("Date", "[DD Month YYYY]"), ("To", "[Name, Title]"),
                    ("From", "[Name, Title]"), ("CC", "[Names]"), ("Subject", "[Subject]")])
    P(d, "Summary", style="Heading 2")
    P(d, ph="[One-paragraph summary of the matter.]")
    P(d, "Details", style="Heading 2")
    for i in range(3):
        P(d, style="List Number", ph=f"[Point {i + 1}]")
    callout(d, "Action required", [("[What is required, by whom, and by when]",)])
    P(d, "Approval", style="Heading 2")
    signature_block(d, ["Prepared by", "Approved by"])
    return d


# ====================================================================== report
def report():
    d = new(margins=(28, 22, 26, 24), title="Report template", subject="Destination Nigeria report")
    s = d.sections[0]
    # --- cover (own section, no header/footer)
    cp = d.paragraphs[0] if d.paragraphs else d.add_paragraph()
    float_picture(cp, ASSETS / "cover_report.jpg", 0, 0, 210, 168)
    float_picture(cp, LOGO, 210 - 24 - 46, 297 - 24 - 46 * LOGO_H / LOGO_W, 46, 46 * LOGO_H / LOGO_W, behind=False)
    for _ in range(18):
        P(d, "")
    q = P(d, style="Small")
    run(q, PHOTOS["zuma-rock"]["caption"])
    P(d, "")
    P(d, ph="[Report type — e.g. Tourism Development Report]", style="Overline")
    q = P(d, style="Title")
    field(q, "[Report title on up to two lines]", "title", color=C["forest"], size=34)
    q = P(d, style="Subtitle")
    field(q, "[Subtitle or reporting period]", "subtitle", color=C["stone"], size=13)
    q = P(d, style="Small")
    field(q, "[Month YYYY]  ·  [Version 1.0]  ·  [Classification: Public / Internal]", "meta")
    # --- section 2: front matter
    s2 = d.add_section(WD_SECTION.NEW_PAGE)
    small_header(s2, "[Report title]")
    contact_footer(s2, "[Report title]  ·  [Month YYYY]")
    s.different_first_page_header_footer = False
    for part in (s.header, s.footer):
        part.is_linked_to_previous = False
    P(d, "Contents", style="Title")
    toc(d, [(1, "Executive summary", 3), (1, "1  Context", 5), (2, "1.1  Background", 5), (2, "1.2  Scope and method", 5),
            (1, "2  Findings", 5), (2, "2.1  Key data", 5), (2, "2.2  Stakeholder perspectives", 6),
            (1, "3  Recommendations", 6), (1, "References", 6), (1, "Appendices", 7)])
    P(d, "To refresh: right-click the table → Update Field → Update entire table.", style="Caption")
    page_break(d)
    P(d, "Summary", style="Overline")
    P(d, "Executive summary", style="Title")
    P(d, ph="[Lead paragraph — the single most important conclusion of the report in two or three sentences.]", style="Lead")
    P(d, ph="[Body — summarise the purpose, approach, main findings and recommendations.]")
    data_table(d, ["Indicator", "Baseline", "Current", "Source"],
               [[("[Indicator]",), ("[Value]",), ("[Value]",), ("[Source, date]",)] for _ in range(3)], [62, 30, 30, 44], align_right=(1, 2))
    callout(d, "Key recommendation", [("[One-sentence recommendation]",)])
    # --- chapter divider
    s3 = d.add_section(WD_SECTION.NEW_PAGE)
    s3.different_first_page_header_footer = True
    for part in (s3.first_page_header, s3.first_page_footer):
        part.is_linked_to_previous = False
    dp = d.add_paragraph()
    float_picture(dp, ASSETS / "divider.jpg", 0, 0, 210, 297)
    for _ in range(14):
        P(d, "")
    run(P(d, ""), "01", font=DISPLAY, size=72, color=C["sage"])
    q = P(d, style="Title")
    run(q, "Context", color=C["ivory"])
    q = P(d, style="Lead")
    field(q, "[Short chapter introduction — one sentence]", color=C["sage"])
    # --- body
    s4 = d.add_section(WD_SECTION.NEW_PAGE)
    s4.different_first_page_header_footer = False
    small_header(s4, "[Report title]")
    contact_footer(s4, "[Report title]  ·  [Month YYYY]")
    P(d, "Chapter 1", style="Overline")
    P(d, "1  Context", style="Heading 1")
    P(d, "1.1  Background", style="Heading 2")
    P(d, ph="[Body text. Public Sans 10 pt on 13 pt leading. Keep paragraphs short. Cite sources inline as (Author, Year) and list them under References.]")
    P(d, ph="[Second paragraph.]")
    P(d, "1.2  Scope and method", style="Heading 2")
    for t in ["[Method or data source 1]", "[Method or data source 2]", "[Limitations]"]:
        P(d, style="List Bullet", ph=t)
    p, r = picture(d, ASSETS / "photo_body.jpg", 166)
    q = P(d, style="Caption")
    run(q, "Figure 1  " + PHOTOS["idanre-hills"]["caption"] + ".  ")
    field(q, "Photo: [photographer and licence to be confirmed]")
    P(d, "2  Findings", style="Heading 1")
    P(d, "2.1  Key data", style="Heading 2")
    data_table(d, ["Destination / State", "Category", "Status", "Notes"],
               [[("[Name]",), ("[Nature / Culture / …]",), ("[Status]",), ("[Notes]",)] for _ in range(5)], [48, 38, 30, 50])
    picture(d, ASSETS / "chart.png", 166)
    q = P(d, style="Caption")
    run(q, "Figure 2  ")
    field(q, "[Chart title]. Source: [verified source, year]. Replace the illustrative chart with the real data.")
    q = P(d, style="Quote")
    field(q, "“[Pull quote from a stakeholder — attributed and approved.]”", color=C["forest"], size=16)
    q = P(d, style="Small")
    field(q, "— [Name, role, organisation]")
    P(d, "2.2  Stakeholder perspectives", style="Heading 2")
    P(d, ph="[Summary of consultation findings.]")
    P(d, "3  Recommendations", style="Heading 1")
    for i in range(1, 4):
        q = P(d, "")
        run(q, f"{i:02d}  ", bold=True, color=C["golddeep"], size=12)
        field(q, f"[Recommendation {i} — action, owner, timeframe]", color=C["charcoal"], size=10.5)
    P(d, "References", style="Heading 1")
    for i in range(3):
        P(d, style="Small", ph="[Author (Year). Title. Publisher. URL, accessed DD Month YYYY.]")
    P(d, "Appendices", style="Heading 1")
    P(d, "Appendix A  —  ", style="Heading 3", ph="[Title]")
    P(d, ph="[Appendix content.]")
    return d


# ====================================================================== proposal / concept note
def proposal():
    d = new(title="Proposal / concept note", subject="Partnership and programme proposals")
    s = d.sections[0]
    s.different_first_page_header_footer = True
    fh = s.first_page_header
    fh.is_linked_to_previous = False
    cp = d.paragraphs[0] if d.paragraphs else d.add_paragraph()
    float_picture(cp, ASSETS / "cover_proposal.jpg", 0, 0, 210, 297)
    float_picture(cp, LOGO, 210 - 24 - 46, 297 - 24 - 46 * LOGO_H / LOGO_W, 46, 46 * LOGO_H / LOGO_W, behind=False)
    for _ in range(20):
        P(d, "")
    q = P(d, style="Overline")
    field(q, "[Proposal type — Partnership / Sponsorship / Programme concept / Project]", color=C["golddeep"], bold=True)
    q = P(d, style="Title")
    field(q, "[Proposal title]", color=C["forest"], size=34)
    q = P(d, style="Lead")
    field(q, "[Prepared for: Organisation]", color=C["stone"])
    q = P(d, style="Small")
    field(q, "[Month YYYY]  ·  [Reference]  ·  [Confidential / For discussion]", color=C["stone"])
    page_break(d)
    small_header(s, "[Proposal title]")
    contact_footer(s, "[Proposal title]")
    secs = [("Summary", "[What is proposed, for whom, and the expected result — 80 words.]"),
            ("Objectives", None), ("Background", "[Context and need, with sources.]"),
            ("Strategic alignment", "[How the proposal supports the Destination Nigeria pillars (Experience · Invest · Explore) and any partner priorities — cite approved documents only.]")]
    for i, (h, body) in enumerate(secs, 1):
        P(d, f"{i:02d}", style="Overline")
        P(d, h, style="Heading 1")
        if body:
            P(d, ph=body)
        else:
            for k in range(3):
                P(d, style="List Number", ph=f"[Objective {k + 1} — specific and measurable]")
    P(d, "05", style="Overline")
    P(d, "Implementation plan", style="Heading 1")
    data_table(d, ["Phase", "Activities", "Lead", "Timing"], [[f"{k}", ("[Activities]",), ("[Owner]",), ("[Month–Month YYYY]",)] for k in range(1, 4)], [18, 80, 34, 34])
    P(d, "06", style="Overline")
    P(d, "Outcomes and indicators", style="Heading 1")
    data_table(d, ["Outcome", "Indicator", "Target", "Means of verification"], [[("[Outcome]",), ("[Indicator]",), ("[Target]",), ("[Source]",)] for _ in range(3)], [44, 44, 30, 48])
    P(d, "07", style="Overline")
    P(d, "Timeline", style="Heading 1")
    t = data_table(d, ["Task", "M1", "M2", "M3", "M4", "M5", "M6"], [[("[Workstream]",), "", "", "", "", "", ""] for _ in range(4)], [52, 19, 19, 19, 19, 19, 19])
    for r_i, row in enumerate(t.rows[1:]):
        for c_i in range(1 + r_i, 3 + r_i):
            cell_shade(row.cells[c_i], C["forest"] if r_i % 2 == 0 else C["sage"])
    P(d, "08", style="Overline")
    P(d, "Roles and responsibilities", style="Heading 1")
    data_table(d, ["Party", "Responsibilities"], [[("[Destination Nigeria]",), ("[Responsibilities]",)], [("[Partner]",), ("[Responsibilities]",)]], [50, 116])
    P(d, "09", style="Overline")
    P(d, "Budget", style="Heading 1")
    data_table(d, ["Item", "Description", "Qty", "Unit cost (₦)", "Total (₦)"],
               [[("[Item]",), ("[Description]",), ("[0]",), ("[0.00]",), ("[0.00]",)] for _ in range(4)] + [["Total", "", "", "", ("[0.00]",)]],
               [34, 62, 14, 28, 28], align_right=(2, 3, 4), total_row=True)
    P(d, "All figures are estimates until confirmed in writing. Taxes and fees per applicable regulations.", style="Caption")
    P(d, "10", style="Overline")
    P(d, "Risks and mitigation", style="Heading 1")
    data_table(d, ["Risk", "Likelihood", "Impact", "Mitigation"], [[("[Risk]",), ("[L/M/H]",), ("[L/M/H]",), ("[Mitigation]",)] for _ in range(3)], [56, 24, 24, 62])
    P(d, "11", style="Overline")
    P(d, "Next steps", style="Heading 1")
    for k in range(3):
        q = P(d, "")
        run(q, f"{k + 1:02d}  ", bold=True, color=C["golddeep"])
        field(q, f"[Step {k + 1} — owner, date]")
    callout(d, "Contact", [("[Name, title]",), f'{_CONTACT["email"]}   {_CONTACT["phone"]}'])
    return d


# ====================================================================== minutes
def minutes():
    d = new(title="Meeting minutes")
    s = d.sections[0]
    logo_header(s, 40)
    contact_footer(s, "Minutes of meeting")
    P(d, "Minutes of meeting", style="Overline")
    q = P(d, style="Title")
    field(q, "[Meeting title]", color=C["forest"], size=28)
    field_table(d, [("Date", "[DD Month YYYY]"), ("Time", "[HH:MM – HH:MM]"), ("Venue", "[Venue / virtual platform]"),
                    ("Chair", "[Name]"), ("Minutes by", "[Name]")])
    P(d, "Participants", style="Heading 2")
    data_table(d, ["Name", "Organisation", "Role", "Present"], [[("[Name]",), ("[Organisation]",), ("[Role]",), "☐"] for _ in range(5)], [48, 52, 46, 20])
    P(d, "Apologies: ", style="Small", ph="[Names]")
    P(d, "Agenda", style="Heading 2")
    for k in range(4):
        P(d, style="List Number", ph=f"[Agenda item {k + 1}]")
    P(d, "Discussion summary", style="Heading 2")
    for k in range(1, 3):
        P(d, f"Item {k}", style="Heading 3")
        P(d, ph="[Summary of discussion — key points, not a transcript.]")
    P(d, "Decisions", style="Heading 2")
    data_table(d, ["No.", "Decision"], [[f"D{k}", ("[Decision]",)] for k in range(1, 4)], [16, 150])
    P(d, "Action items", style="Heading 2")
    data_table(d, ["No.", "Action", "Owner", "Deadline", "Status"],
               [[f"A{k}", ("[Action]",), ("[Owner]",), ("[DD/MM/YYYY]",), ("[Open]",)] for k in range(1, 5)], [14, 72, 30, 28, 22])
    callout(d, "Next meeting", [("[Date · Time · Venue]",)])
    P(d, "Minutes approved", style="Heading 2")
    signature_block(d, ["Chair", "Secretary"])
    return d


# ====================================================================== invoice & quotation
def invoice(kind="Invoice"):
    d = new(title=kind)
    s = d.sections[0]
    logo_header(s, 44)
    contact_footer(s, f"{kind}")
    P(d, kind, style="Title")
    t = d.add_table(rows=1, cols=2)
    borders(t, "#FFFFFF", 0, outer=False)
    a, b = t.rows[0].cells
    for cell, title, fields in [(a, "From", ["Destination Nigeria", "[Registered business name and address — verified]", f'{_CONTACT["email"]} · {_CONTACT["phone"]}', "[Tax ID (TIN) — if applicable]"]),
                                (b, "Bill to (client)" if kind == "Invoice" else "Prepared for", ["[Client name]", "[Organisation]", "[Address]", "[Client email]"])]:
        p = cell.paragraphs[0]
        p.style = d.styles["Overline"]
        run(p, title)
        for f in fields:
            q = cell.add_paragraph(style="Table Text")
            field(q, f)
    widths(t, [83, 83])
    P(d, "", style="Small")
    rows = [("Number", "[INV-YYYY-0001]" if kind == "Invoice" else "[QUO-YYYY-0001]"), ("Issue date", "[DD Month YYYY]"),
            ("Due date" if kind == "Invoice" else "Valid until", "[DD Month YYYY]"), ("Reference / PO", "[Reference]")]
    field_table(d, rows, 40, 126)
    data_table(d, ["Description of service", "Qty", "Unit price (₦)", "Amount (₦)"],
               [[("[Service description]",), ("[0]",), ("[0.00]",), ("[0.00]",)] for _ in range(5)], [86, 16, 32, 32], align_right=(1, 2, 3), zebra=C["ivory"])
    t2 = d.add_table(rows=0, cols=2)
    borders(t2, C["sage"], 4)
    cell_margins(t2)
    for label, ph, bold in [("Subtotal", "[0.00]", False), ("VAT ([rate]% — if applicable)", "[0.00]", False), ("Other applicable taxes / levies", "[0.00]", False), ("Total (₦)", "[0.00]", True)]:
        x, y = t2.add_row().cells
        px = x.paragraphs[0]
        px.style = d.styles["Table Head"] if bold else d.styles["Table Text"]
        run(px, label)
        py = y.paragraphs[0]
        py.style = d.styles["Table Text"]
        py.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        field(py, ph, bold=bold, color=C["forest"] if bold else C["stone"])
        if bold:
            rule_cell(x, "top", C["forest"], 8)
            rule_cell(y, "top", C["forest"], 8)
    widths(t2, [134, 32])
    P(d, "")
    if kind == "Invoice":
        callout(d, "Payment terms", [("[e.g. Payment due within 30 days of invoice date]",),
                                     ("[Payment details are issued separately on official letterhead and must be verified by phone before transfer.]",)])
        P(d, "Do not pay into any account that is not confirmed directly with Destination Nigeria’s finance contact.", style="Caption")
    else:
        callout(d, "Terms", [("[Scope, assumptions, exclusions, validity period and acceptance process]",)])
        signature_block(d, ["Accepted for the client", "For Destination Nigeria"])
    return d


# ====================================================================== press & media
def press(kind):
    d = new(title=kind)
    s = d.sections[0]
    logo_header(s, 44)
    contact_footer(s, kind)
    labels = {"Press release": "For immediate release", "Media advisory": "Media advisory", "Official announcement": "Official announcement",
              "Interview brief": "Interview brief — not for publication", "Fact sheet": "Fact sheet", "Media backgrounder": "Media backgrounder"}
    q = P(d, style="Overline")
    run(q, labels[kind])
    if kind == "Press release":
        run(q, "  ·  ")
        field(q, "[or EMBARGOED UNTIL HH:MM, DD Month YYYY]", color=C["golddeep"], bold=True)
    q = P(d, style="Title")
    field(q, "[Headline — active verb, under 12 words]" if kind != "Interview brief" else "[Interview: Name — Outlet]", color=C["forest"], size=26)
    if kind in ("Press release", "Official announcement"):
        q = P(d, style="Subtitle")
        field(q, "[Sub-headline — the key supporting fact]", color=C["stone"], size=13)
        q = P(d, "")
        run(q, "", bold=True)
        field(q, "[CITY], [DD Month YYYY]", bold=True, color=C["forest"])
        run(q, " — ")
        field(q, "[Lead paragraph: who, what, when, where, why — verified facts only.]")
        P(d, ph="[Second paragraph: context and significance.]")
        q = P(d, style="Quote")
        field(q, "“[Approved quotation.]”", color=C["forest"], size=15)
        P(d, style="Small", ph="— [Name, Title, Organisation]")
        P(d, ph="[Further detail, partners (approved only), availability of images/interviews.]")
        P(d, "— ENDS —", align="center", bold=True, color=C["forest"])
    elif kind == "Media advisory":
        field_table(d, [("What", "[Event / activity]"), ("When", "[Date, time — confirmed]"), ("Where", "[Venue, address]"),
                        ("Who", "[Speakers — confirmed only]"), ("Why", "[Purpose / news value]"), ("Media access", "[Accreditation, arrival time, photo/interview opportunities]"),
                        ("RSVP", "[Name · email · deadline]")])
    elif kind == "Interview brief":
        field_table(d, [("Outlet / format", "[Outlet · live/recorded · length]"), ("Date / time", "[DD Month YYYY · HH:MM]"), ("Interviewer", "[Name]"),
                        ("Spokesperson", "[Name, title]"), ("Audience", "[Audience]")])
        P(d, "Key messages", style="Heading 2")
        for m in ["Experience: [message]", "Invest: [message]", "Explore: [message]"]:
            P(d, style="List Number", ph=f"[{m}]")
        P(d, "Likely questions and approved answers", style="Heading 2")
        data_table(d, ["Question", "Approved answer / bridge"], [[("[Question]",), ("[Answer]",)] for _ in range(4)], [70, 96])
        P(d, "Facts to use (verified, with source)", style="Heading 2")
        data_table(d, ["Fact", "Source and date"], [[("[Fact]",), ("[Source]",)] for _ in range(3)], [100, 66])
        callout(d, "Avoid", [("[Topics or claims that are not confirmed]",)])
    elif kind == "Fact sheet":
        P(d, ph="[One-sentence description of the subject.]", style="Lead")
        t = d.add_table(rows=1, cols=3)
        borders(t, "#FFFFFF", 0, outer=False)
        for c in t.rows[0].cells:
            rule_cell(c, "top", C["forest"], 8)
            p = c.paragraphs[0]
            field(p, "[Figure]", color=C["forest"], bold=False, size=30)
            for r_ in p.runs:
                r_.font.name = DISPLAY
            q = c.add_paragraph(style="Small")
            field(q, "[What the figure measures — source, year]", color=C["stone"])
        cell_margins(t, 120, 120, 0, 160)
        P(d, "")
        P(d, "At a glance", style="Heading 2")
        field_table(d, [("Overview", "[Text]"), ("Location", "[Text]"), ("Highlights", "[Text]"), ("Getting there", "[Text]"), ("Best time to visit", "[Text]"), ("Further information", "[URL]")])
        P(d, "All figures must cite a verifiable source. Remove any figure that cannot be verified.", style="Caption")
    elif kind == "Media backgrounder":
        P(d, ph="[Standfirst — what Destination Nigeria is, in two sentences.]", style="Lead")
        for h in ["About the platform", "The proposition: Experience. Invest. Explore.", "How it works", "Partners and stakeholders (approved list only)", "Key dates (confirmed only)"]:
            P(d, h, style="Heading 2")
            P(d, ph="[Text]")
    if kind not in ("Interview brief",):
        P(d, "About Destination Nigeria", style="Heading 3")
        P(d, ph="[Approved boilerplate — 60–80 words. Keep identical across all releases.]", style="Small")
        P(d, "Media contact", style="Heading 3")
        P(d, ph="[Name · Title · Email · Telephone]", style="Small")
        q = P(d, style="Small")
        run(q, "Source verification: ", bold=True, color=C["forest"])
        run(q, "every fact, figure, name and date in this document has been checked against [source] by [name] on [date]. Delete this line before issue.", color=C["stone"])
    return d


# ====================================================================== workshop materials
def workshop(kind):
    land = kind in ("Attendance sheet", "Participant certificate", "Action-plan tracker")
    d = new(landscape=land, margins=(26, 18, 20, 18) if land else (30, 22, 24, 22), title=kind)
    s = d.sections[0]
    if kind == "Participant certificate":
        s.top_margin = Mm(20)
        cp = d.paragraphs[0] if d.paragraphs else d.add_paragraph()
        float_picture(cp, ASSETS / "certificate_bg.jpg", 0, 0, 297, 210)
        p = P(d, "", align="center")
        p.add_run().add_picture(str(LOGO), width=Mm(64))
        P(d, "Certificate of participation", style="Overline", align="center")
        P(d, "This is to certify that", align="center", color=C["stone"])
        q = P(d, style="Title", align="center")
        field(q, "[Participant full name]", color=C["forest"], size=36)
        P(d, "participated in", align="center", color=C["stone"])
        q = P(d, style="Lead", align="center")
        field(q, "[Workshop / programme title]", color=C["forest"], bold=True, size=15)
        q = P(d, align="center")
        field(q, "[Venue, City]  ·  [DD–DD Month YYYY]")
        t = d.add_table(rows=1, cols=2)
        borders(t, "#FFFFFF", 0, outer=False)
        for c, role in zip(t.rows[0].cells, ["[Name, title]", "[Name, title]"]):
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(20)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bottom_rule(p, C["forest"], 6, 2)
            q = c.add_paragraph(style="Small")
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            field(q, role)
        widths(t, [120, 120])
        q = P(d, style="Small", align="center")
        run(q, "Certificate no. ")
        field(q, "[DN-CERT-YYYY-0000]")
        return d
    logo_header(s, 38)
    contact_footer(s, kind)
    q = P(d, style="Overline")
    field(q, "[Workshop title]", color=C["golddeep"], bold=True)
    P(d, kind, style="Title")
    if kind == "Workshop agenda":
        field_table(d, [("Date", "[DD Month YYYY]"), ("Venue", "[Venue]"), ("Facilitator", "[Name]"), ("Objective", "[Objective]")])
        data_table(d, ["Time", "Session", "Lead", "Format"], [[("[09:00]",), ("[Session title]",), ("[Name]",), ("[Plenary / Group]",)] for _ in range(8)], [22, 82, 34, 28], zebra=C["ivory"])
    elif kind == "Attendance sheet":
        field_table(d, [("Date", "[DD Month YYYY]"), ("Venue", "[Venue]")], 30, 230)
        data_table(d, ["No.", "Full name", "Organisation", "Position", "Email", "Telephone", "Signature"],
                   [[str(k), "", "", "", "", "", ""] for k in range(1, 13)], [12, 52, 48, 36, 52, 32, 29])
        P(d, "Personal data on this sheet is used only for attendance records and follow-up for this programme. [Insert approved data-protection notice.]", style="Caption")
    elif kind == "Feedback form":
        P(d, "Thank you for taking part. Your feedback improves future programmes. Responses are anonymous unless you add your name.", style="Lead")
        for qn_, opts in [("Overall, how would you rate the workshop?", ["1", "2", "3", "4", "5"]),
                          ("How relevant was the content to your work?", ["1", "2", "3", "4", "5"]),
                          ("How would you rate the facilitation?", ["1", "2", "3", "4", "5"]),
                          ("Was the venue and organisation satisfactory?", ["Yes", "Partly", "No"])]:
            P(d, qn_, style="Heading 3")
            checkbox_row(P(d, ""), opts)
        for qn_ in ["What was most useful?", "What should we improve?", "What support would you need next?"]:
            P(d, qn_, style="Heading 3")
            for _ in range(3):
                p = P(d, "")
                bottom_rule(p, C["sage"], 4, 1)
        P(d, "Name (optional): ____________________     Organisation (optional): ____________________", style="Small")
    elif kind == "Workshop report":
        field_table(d, [("Date / venue", "[Text]"), ("Participants", "[Number and organisations]"), ("Facilitators", "[Names]")])
        for h in ["Purpose", "Sessions and key points", "Outcomes", "Agreed actions", "Feedback summary", "Recommendations"]:
            P(d, h, style="Heading 2")
            P(d, ph="[Text]")
        picture(d, ASSETS / "photo_body.jpg", 166)
        P(d, style="Caption", ph="[Caption · Photo: name / licence · consent obtained: yes/no]")
    elif kind == "Action-plan tracker":
        data_table(d, ["No.", "Action", "Owner", "Partners", "Start", "Due", "Status", "Notes"],
                   [[str(k), ("[Action]",), ("[Owner]",), ("[Partners]",), ("[DD/MM]",), ("[DD/MM]",), ("[Not started]",), ""] for k in range(1, 11)],
                   [10, 70, 30, 34, 18, 18, 26, 55])
        P(d, "Status key:  Not started  ·  In progress  ·  Done  ·  Blocked", style="Caption")
    return d


# ====================================================================== investment opportunity profile
def investment():
    d = new(title="Investment opportunity profile", margins=(28, 22, 24, 22))
    s = d.sections[0]
    small_header(s, "Investment opportunity profile")
    contact_footer(s, "Investment opportunity profile  ·  [Profile ID]")
    p = d.add_paragraph()
    float_picture(p, ASSETS / "photo_invest.jpg", 0, 0, 210, 92)
    for _ in range(9):
        P(d, "")
    P(d, "Invest  ·  Investment opportunity profile", style="Overline")
    q = P(d, style="Title")
    field(q, "[Project title]", color=C["forest"], size=30)
    q = P(d, style="Lead")
    field(q, "[Location — State, LGA]  ·  [Opportunity category]")
    callout(d, "Data status", ["Only verified project data appears in this profile. Unconfirmed values are shown as placeholders and must not be published.",
                               ("[Verified by: name, role · date]",)])
    field_table(d, [("Opportunity category", "[e.g. Hospitality / Eco-tourism / Heritage / Events infrastructure]"),
                    ("Location", "[State, LGA, coordinates — verified]"), ("Project description", "[Description]"),
                    ("Current development stage", "[Concept / Feasibility / Design / Permitting / Construction / Operational]"),
                    ("Infrastructure context", "[Road, air, power, water, telecoms — verified]"),
                    ("Indicative investment", "[TBC — do not publish unconfirmed values]"),
                    ("Commercial structure", "[TBC — e.g. PPP, concession, JV, equity]"),
                    ("Land / title status", "[Verified status or ‘Under review’]"),
                    ("Supporting documentation", "[Feasibility study, EIA, title, masterplan — list available on request]")], 52, 114)
    P(d, "Due-diligence notes", style="Heading 2")
    data_table(d, ["Item", "Status", "Notes"], [[("[Item]",), ("[Verified / Pending / Not available]",), ("[Notes]",)] for _ in range(5)], [56, 44, 66])
    P(d, "Contact", style="Heading 2")
    P(d, ph="[Name · Title · Email · Telephone]")
    P(d, "This profile is for information only and does not constitute an offer, commitment or financial advice.", style="Caption")
    return d


# ====================================================================== spreadsheets (calculating)
def spreadsheets():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XImg
    thin = Side(style="thin", color=C["sage"].lstrip("#"))
    for kind in ("Invoice", "Quotation"):
        wb = Workbook()
        ws = wb.active
        ws.title = kind
        ws.sheet_view.showGridLines = False
        for col, w in zip("ABCDE", [6, 52, 10, 18, 18]):
            ws.column_dimensions[col].width = w
        img = XImg(str(KIT / "02_LOGO_MASTER_FILES/PNG/DN_logo_primary_fullcolour_500px.png"))
        img.width, img.height = 200, round(200 * LOGO_H / LOGO_W)
        ws.add_image(img, "B1")
        ws["D2"] = kind.upper()
        ws["D2"].font = Font(name="Newsreader Display", size=22, color=C["forest"].lstrip("#"))
        rows = [("Number", "[INV-YYYY-0001]" if kind == "Invoice" else "[QUO-YYYY-0001]"), ("Issue date", "[DD/MM/YYYY]"),
                ("Due date" if kind == "Invoice" else "Valid until", "[DD/MM/YYYY]"), ("Client", "[Client name]"), ("Reference / PO", "[Reference]")]
        for i, (a, b) in enumerate(rows, 6):
            ws[f"B{i}"], ws[f"C{i}"] = a, b
            ws[f"B{i}"].font = Font(name="Public Sans", bold=True, size=9, color=C["forest"].lstrip("#"))
            ws[f"C{i}"].font = Font(name="Public Sans", size=10, color=C["stone"].lstrip("#"))
        hdr = 13
        for c, h in zip("ABCDE", ["No.", "Description of service", "Qty", "Unit price (₦)", "Amount (₦)"]):
            cell = ws[f"{c}{hdr}"]
            cell.value = h
            cell.font = Font(name="Public Sans", bold=True, size=9, color=C["forest"].lstrip("#"))
            cell.fill = PatternFill("solid", fgColor=C["haze"].lstrip("#"))
            cell.border = Border(bottom=thin)
        for r in range(hdr + 1, hdr + 11):
            ws[f"A{r}"] = r - hdr
            ws[f"E{r}"] = f'=IF(C{r}="","",C{r}*D{r})'
            for c in "ABCDE":
                ws[f"{c}{r}"].font = Font(name="Public Sans", size=10)
                ws[f"{c}{r}"].border = Border(bottom=thin)
            for c in "DE":
                ws[f"{c}{r}"].number_format = '#,##0.00'
        last = hdr + 10
        ws[f"D{last + 2}"], ws[f"E{last + 2}"] = "Subtotal", f"=SUM(E{hdr + 1}:E{last})"
        ws[f"C{last + 3}"], ws[f"D{last + 3}"], ws[f"E{last + 3}"] = 0, "VAT rate (enter)", f"=E{last + 2}*C{last + 3}"
        ws[f"C{last + 3}"].number_format = "0.0%"
        ws[f"D{last + 4}"], ws[f"E{last + 4}"] = "Other taxes / levies", 0
        ws[f"D{last + 5}"], ws[f"E{last + 5}"] = "TOTAL (₦)", f"=E{last + 2}+E{last + 3}+E{last + 4}"
        for r in range(last + 2, last + 6):
            ws[f"E{r}"].number_format = '#,##0.00'
            ws[f"D{r}"].font = Font(name="Public Sans", bold=r == last + 5, size=10, color=C["forest"].lstrip("#"))
            ws[f"E{r}"].font = Font(name="Public Sans", bold=r == last + 5, size=11 if r == last + 5 else 10)
        ws[f"D{last + 5}"].fill = ws[f"E{last + 5}"].fill = PatternFill("solid", fgColor=C["haze"].lstrip("#"))
        ws[f"B{last + 8}"] = "Payment / acceptance terms: [enter]. Payment details are issued separately on official letterhead and must be verified by phone."
        ws[f"B{last + 8}"].font = Font(name="Public Sans", size=8, color=C["stone"].lstrip("#"))
        ws[f"B{last + 8}"].alignment = Alignment(wrap_text=True)
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        p = D / "INVOICES" / f"DN_{kind}_Template_A4.xlsx"
        wb.save(p)
        saved.append(p)
    wb = Workbook()
    ws = wb.active
    ws.title = "Action plan"
    heads = ["No.", "Action", "Owner", "Partners", "Start", "Due", "Status", "Notes"]
    for i, (h, w) in enumerate(zip(heads, [6, 50, 20, 22, 12, 12, 14, 40]), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name="Public Sans", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C["forest"].lstrip("#"))
        ws.column_dimensions[chr(64 + i)].width = w
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Not started,In progress,Done,Blocked"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(2, 42):
        ws.cell(row=r, column=1, value=r - 1)
        dv.add(f"G{r}")
        for c in (5, 6):
            ws.cell(row=r, column=c).number_format = "DD/MM/YYYY"
    from openpyxl.formatting.rule import CellIsRule
    for val, col in [("Done", C["sage"]), ("Blocked", "E7DFCF"), ("In progress", "F5F0E6")]:
        ws.conditional_formatting.add("G2:G41", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=col.lstrip("#"))))
    ws.freeze_panes = "A2"
    p = D / "WORKSHOP_MATERIALS" / "DN_Action-Plan-Tracker.xlsx"
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    saved.append(p)


if __name__ == "__main__":
    ASSETS = make_assets()
    save(letterhead("standard"), "LETTERHEAD/DN_Letterhead_Standard_A4.docx")
    save(letterhead("ceremonial"), "LETTERHEAD/DN_Letterhead_Ceremonial_A4.docx")
    save(letterhead("standard", sample=True), "LETTERHEAD/DN_Letter_Formal-Correspondence_A4.docx")
    save(memo(), "MEMO/DN_Memo_A4.docx")
    save(report(), "REPORTS/DN_Report_Template_A4.docx")
    save(proposal(), "PROPOSALS/DN_Proposal-Concept-Note_Template_A4.docx")
    save(minutes(), "MEETING_MINUTES/DN_Meeting-Minutes_A4.docx")
    save(invoice("Invoice"), "INVOICES/DN_Invoice_Template_A4.docx")
    save(invoice("Quotation"), "INVOICES/DN_Quotation_Template_A4.docx")
    for k in ["Press release", "Media advisory", "Official announcement", "Interview brief", "Fact sheet", "Media backgrounder"]:
        save(press(k), f"PRESS_AND_MEDIA/DN_{k.title().replace(' ', '-')}_A4.docx")
    for k in ["Workshop agenda", "Attendance sheet", "Participant certificate", "Feedback form", "Workshop report", "Action-plan tracker"]:
        save(workshop(k), f"WORKSHOP_MATERIALS/DN_{k.title().replace(' ', '-')}_A4.docx")
    save(investment(), "INVESTMENT_PROFILE/DN_Investment-Opportunity-Profile_A4.docx")
    spreadsheets()
    print(len(saved), "files")
    json.dump([str(p.relative_to(KIT)) for p in saved], open(KIT / "16_SOURCE_FILES/build/_jobs/_documents.json", "w"), indent=1)
